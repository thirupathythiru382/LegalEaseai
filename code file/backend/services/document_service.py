from io import BytesIO
from pathlib import Path
from typing import Optional
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF


# =========================================================
# PROJECT PATHS
# =========================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

ASSETS_DIR = PROJECT_ROOT / "assets"

LOGO_PATH = ASSETS_DIR / "logo.png"


# =========================================================
# TXT
# =========================================================

def format_txt(text: str) -> bytes:
    return text.encode("utf-8")


# =========================================================
# DOCX FONT
# =========================================================

def set_document_font(document: Document) -> None:

    styles = document.styles

    for style_name in [
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
    ]:

        try:

            style = styles[style_name]

            style.font.name = "Times New Roman"
            style.font.size = Pt(12)

            rpr = style._element.get_or_add_rPr()

            rfonts = rpr.rFonts

            if rfonts is not None:

                rfonts.set(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii",
                    "Times New Roman",
                )

                rfonts.set(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi",
                    "Times New Roman",
                )

        except Exception:
            pass


# =========================================================
# DOCX LOGO
# =========================================================

def add_logo(document: Document) -> None:

    if not LOGO_PATH.exists():
        return

    try:

        section = document.sections[0]

        header = section.header

        paragraph = header.paragraphs[0]

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run()

        run.add_picture(
            str(LOGO_PATH),
            width=Inches(1.25),
        )

    except Exception:
        pass


# =========================================================
# DOCX FOOTER
# =========================================================

def add_footer(document: Document) -> None:

    for section in document.sections:

        footer = section.footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            "Generated with LegalEase AI"
        )

        run.font.name = "Times New Roman"
        run.font.size = Pt(9)


# =========================================================
# HEADING DETECTION
# =========================================================

def is_heading(line: str) -> bool:

    stripped = line.strip()

    if not stripped:
        return False

    heading_keywords = [
        "AGREEMENT",
        "CONTRACT",
        "PARTIES",
        "RECITALS",
        "DEFINITIONS",
        "CONFIDENTIALITY",
        "OBLIGATIONS",
        "PAYMENT",
        "TERM",
        "TERMINATION",
        "GOVERNING LAW",
        "DISPUTE",
        "SIGNATURE",
        "SIGNATURES",
        "GENERAL PROVISIONS",
        "NOTICE",
        "INTELLECTUAL PROPERTY",
        "LIABILITY",
        "INDEMNIFICATION",
    ]

    upper_line = stripped.upper()

    for keyword in heading_keywords:

        if upper_line == keyword:
            return True

        if upper_line.startswith(keyword + ":"):
            return True

    return False


# =========================================================
# DOCX TERMS TABLE
# =========================================================

def add_terms_table(
    document: Document,
    terms: list[str],
) -> None:

    if not terms:
        return

    document.add_heading(
        "Key Terms",
        level=2,
    )

    table = document.add_table(
        rows=1,
        cols=2,
    )

    table.style = "Table Grid"

    header_cells = table.rows[0].cells

    header_cells[0].text = "No."
    header_cells[1].text = "Term"

    for index, term in enumerate(
        terms,
        start=1,
    ):

        cells = table.add_row().cells

        cells[0].text = str(index)
        cells[1].text = term

        cells[0].vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )

        cells[1].vertical_alignment = (
            WD_CELL_VERTICAL_ALIGNMENT.CENTER
        )


# =========================================================
# DOCX
# =========================================================

def format_docx(
    text: str,
    title: str = "LegalEase AI Legal Document",
    terms: Optional[list[str]] = None,
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    set_document_font(document)

    add_logo(document)

    # -----------------------------------------------------
    # Title
    # -----------------------------------------------------

    title_paragraph = document.add_paragraph()

    title_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title_paragraph.add_run(
        title
    )

    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(18)

    # -----------------------------------------------------
    # Content
    # -----------------------------------------------------

    for line in text.splitlines():

        stripped = line.strip()

        if not stripped:

            document.add_paragraph()

            continue

        if is_heading(stripped):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                stripped
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)

            continue

        if stripped.startswith(
            (
                "- ",
                "• ",
                "* ",
                "● ",
                "▪ ",
                "◦ ",
            )
        ):

            bullet_text = stripped[2:].strip()

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            run = paragraph.add_run(
                bullet_text
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

            continue

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
        )

        paragraph.paragraph_format.space_after = Pt(6)

        run = paragraph.add_run(
            stripped
        )

        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

    if terms:

        add_terms_table(
            document,
            terms,
        )

    add_footer(document)

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


# =========================================================
# PDF CLASS
# =========================================================

class LegalPDF(FPDF):

    def header(self):

        # -------------------------------------------------
        # Logo
        # -------------------------------------------------

        if LOGO_PATH.exists():

            try:

                self.image(
                    str(LOGO_PATH),
                    x=10,
                    y=8,
                    w=18,
                )

            except Exception:
                pass

        # -------------------------------------------------
        # Header
        # -------------------------------------------------

        self.set_font(
            "Times",
            "B",
            12,
        )

        available_width = (
            self.w
            - self.l_margin
            - self.r_margin
        )

        self.cell(
            available_width,
            10,
            "LegalEase AI",
            align="C",
        )

        self.ln(12)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Times",
            "",
            8,
        )

        available_width = (
            self.w
            - self.l_margin
            - self.r_margin
        )

        self.cell(
            available_width,
            10,
            (
                f"Page {self.page_no()} "
                "| Generated with LegalEase AI"
            ),
            align="C",
        )


# =========================================================
# PDF CHARACTER CLEANING
# =========================================================

def pdf_safe_text(text: str) -> str:

    replacements = {

        "•": "-",
        "●": "-",
        "▪": "-",
        "◦": "-",

        "–": "-",
        "—": "-",
        "−": "-",

        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",

        "…": "...",

        "™": "(TM)",
        "©": "(C)",
        "®": "(R)",

        "₹": "INR ",

        "\u00a0": " ",
    }

    for old, new in replacements.items():

        text = text.replace(
            old,
            new,
        )

    # Convert unsupported characters
    text = (
        text
        .encode(
            "latin-1",
            errors="replace",
        )
        .decode("latin-1")
    )

    return text


# =========================================================
# PDF MANUAL WORD WRAPPING
# =========================================================

def wrap_pdf_text(
    pdf: FPDF,
    text: str,
    max_width: float,
) -> list[str]:
    """
    Manually wrap text according to actual PDF width.

    This prevents FPDF from throwing:

        Not enough horizontal space
        to render a single character
    """

    text = pdf_safe_text(text)

    if not text:
        return [""]

    words = text.split(" ")

    lines = []

    current_line = ""

    for word in words:

        # -------------------------------------------------
        # Empty token
        # -------------------------------------------------

        if word == "":
            continue

        # -------------------------------------------------
        # Check normal word
        # -------------------------------------------------

        test_line = (
            word
            if not current_line
            else current_line + " " + word
        )

        width = pdf.get_string_width(
            test_line
        )

        # -------------------------------------------------
        # Word fits
        # -------------------------------------------------

        if width <= max_width:

            current_line = test_line

            continue

        # -------------------------------------------------
        # Current line exists
        # -------------------------------------------------

        if current_line:

            lines.append(
                current_line
            )

            current_line = ""

        # -------------------------------------------------
        # Check individual word
        # -------------------------------------------------

        word_width = pdf.get_string_width(
            word
        )

        if word_width <= max_width:

            current_line = word

            continue

        # -------------------------------------------------
        # VERY LONG WORD
        # -------------------------------------------------

        chunk = ""

        for character in word:

            test_chunk = chunk + character

            character_width = (
                pdf.get_string_width(
                    test_chunk
                )
            )

            if (
                character_width <= max_width
            ):

                chunk = test_chunk

            else:

                if chunk:

                    lines.append(
                        chunk
                    )

                chunk = character

        if chunk:

            current_line = chunk

    if current_line:

        lines.append(
            current_line
        )

    return lines


# =========================================================
# PDF WRITE WRAPPED TEXT
# =========================================================

def pdf_write_wrapped(
    pdf: FPDF,
    text: str,
    line_height: float = 7,
) -> None:
    """
    Safely write wrapped text to the PDF.
    """

    available_width = (
        pdf.w
        - pdf.l_margin
        - pdf.r_margin
    )

    # Keep a small safety margin
    available_width -= 2

    lines = wrap_pdf_text(
        pdf,
        text,
        available_width,
    )

    for line in lines:

        # Empty line
        if not line:

            pdf.ln(line_height)

            continue

        pdf.cell(
            available_width,
            line_height,
            line,
            new_x="LMARGIN",
            new_y="NEXT",
        )


# =========================================================
# PDF FORMAT
# =========================================================

def format_pdf(
    text: str,
    title: str = "LegalEase AI Legal Document",
) -> bytes:

    # -----------------------------------------------------
    # Create PDF
    # -----------------------------------------------------

    pdf = LegalPDF(
        orientation="P",
        unit="mm",
        format="A4",
    )

    # -----------------------------------------------------
    # Margins
    # -----------------------------------------------------

    pdf.set_margins(
        left=15,
        top=25,
        right=15,
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20,
    )

    pdf.add_page()

    # -----------------------------------------------------
    # Title
    # -----------------------------------------------------

    pdf.set_font(
        "Times",
        "B",
        16,
    )

    pdf_write_wrapped(
        pdf,
        title,
        line_height=9,
    )

    pdf.ln(5)

    # -----------------------------------------------------
    # Content
    # -----------------------------------------------------

    for original_line in text.splitlines():

        line = original_line.strip()

        # -------------------------------------------------
        # Empty line
        # -------------------------------------------------

        if not line:

            pdf.ln(5)

            continue

        # -------------------------------------------------
        # Heading
        # -------------------------------------------------

        if is_heading(line):

            pdf.ln(2)

            pdf.set_font(
                "Times",
                "B",
                13,
            )

            pdf_write_wrapped(
                pdf,
                line,
                line_height=8,
            )

            pdf.set_font(
                "Times",
                "",
                12,
            )

            pdf.ln(1)

            continue

        # -------------------------------------------------
        # Bullet
        # -------------------------------------------------

        if line.startswith(
            (
                "- ",
                "• ",
                "* ",
                "● ",
                "▪ ",
                "◦ ",
            )
        ):

            bullet_text = line[2:].strip()

            bullet_text = bullet_text.lstrip(
                "•●▪* -"
            )

            pdf.set_font(
                "Times",
                "",
                12,
            )

            pdf_write_wrapped(
                pdf,
                "- " + bullet_text,
                line_height=7,
            )

            continue

        # -------------------------------------------------
        # Normal paragraph
        # -------------------------------------------------

        pdf.set_font(
            "Times",
            "",
            12,
        )

        pdf_write_wrapped(
            pdf,
            line,
            line_height=7,
        )

        pdf.ln(2)

    # -----------------------------------------------------
    # Output
    # -----------------------------------------------------

    output = pdf.output()

    if isinstance(output, bytes):

        return output

    return bytes(output)


# =========================================================
# ALL FORMATS
# =========================================================

def create_all_formats(
    text: str,
) -> dict[str, bytes]:

    return {
        "txt": format_txt(text),
        "docx": format_docx(text),
        "pdf": format_pdf(text),
    }